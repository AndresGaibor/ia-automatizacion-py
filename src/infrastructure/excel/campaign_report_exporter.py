from typing import Optional, Union
from openpyxl.utils import get_column_letter

from src.core.dto import ScrapingResult
from src.shared.logging.logger import get_logger

logger = get_logger()

CAMPAIGN_REPORT_HEADERS = [
    "Buscar",
    "Nombre",
    "ID Campaña",
    "Fecha",
    "Total enviado",
    "Abierto",
    "No abierto",
    "URL de Correo",
]


class CampaignReportExporter:
    def __init__(self, max_col_width: int = 50, col_padding: int = 2):
        self.max_col_width = max_col_width
        self.col_padding = col_padding

    def export(
        self,
        campaign_data: Union[ScrapingResult, list[list[str]]],
        output_path: str,
        sheet_name: str = "Sheet",
    ) -> None:
        """Exporta datos de campañas a Excel.

        Args:
            campaign_data: Puede ser ScrapingResult o list[list[str]]
                          (para backward compatibility con código legacy)
            output_path: Ruta del archivo Excel de salida
            sheet_name: Nombre de la hoja (default: "Sheet")
        """
        from openpyxl import Workbook, load_workbook

        # Convertir list[list[str]] a filas si no es ScrapingResult
        if isinstance(campaign_data, ScrapingResult):
            rows = campaign_data.to_excel_rows()
            total = campaign_data.total_campaigns
        else:
            rows = campaign_data
            total = len(campaign_data)

        wb = _load_or_create(output_path)
        ws = _get_or_create_sheet(wb, sheet_name)
        _clear_from(ws, 1)

        ws.append(CAMPAIGN_REPORT_HEADERS)

        for row in rows:
            ws.append(row)

        _auto_fit_columns(ws, self.max_col_width, self.col_padding)
        wb.save(output_path)
        logger.success(f"Archivo guardado: {output_path} ({total} registros)")


def _load_or_create(path: str):
    from openpyxl import Workbook, load_workbook
    import os

    try:
        return load_workbook(path)
    except FileNotFoundError:
        return Workbook()


def _get_or_create_sheet(wb, name: str):
    if name in wb.sheetnames:
        return wb[name]
    return wb.create_sheet(title=name)


def _clear_from(ws, start_row: int):
    max_row = ws.max_row
    if max_row >= start_row:
        ws.delete_rows(start_row, max_row - start_row + 1)


def _auto_fit_columns(ws, max_width: int, padding: int):
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(1, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if cell.value and len(str(cell.value)) > max_len:
                max_len = len(str(cell.value))
        ws.column_dimensions[col_letter].width = min(max_len + padding, max_width)


def load_campaign_listing_progress(path: str) -> tuple[list[list[str]], set[str]]:
    from openpyxl import load_workbook
    import os

    if not os.path.exists(path):
        return [], set()

    try:
        wb = load_workbook(path)
        ws = wb.active
        filas = list(ws.iter_rows(values_only=True))
        wb.close()

        if not filas:
            return [], set()

        encabezados = list(filas[0])
        if "URL de Correo" not in encabezados:
            return [], set()

        campanias: list[list[str]] = []
        ids: set[str] = set()

        for fila in filas[1:]:
            valores = ["" if v is None else v for v in list(fila)]
            if len(valores) >= 3 and valores[2]:
                ids.add(str(valores[2]))
                while len(valores) < 8:
                    valores.append("")
                campanias.append(valores)

        return campanias, ids

    except Exception:
        return [], set()


def save_campaign_listing_progress(path: str, campaigns: list[list[str]], headers: list[str]) -> None:
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet"

    ws.append(headers)

    for campania in campaigns:
        fila = list(campania)
        while len(fila) < 8:
            fila.append("")
        ws.append(fila)

    _auto_fit_columns(ws, max_width=50, padding=2)

    wb.save(path)
    wb.close()
